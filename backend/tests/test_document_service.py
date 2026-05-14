from app.models.document import Document
from app.services.document_service import create_document_record, extract_text


def test_create_document_record_sets_processing_status(db_session):
    doc = create_document_record(
        db=db_session,
        filename="notes.txt",
        file_path="/tmp/notes.txt",
    )

    saved = db_session.query(Document).filter(Document.id == doc.id).one()
    assert saved.filename == "notes.txt"
    assert saved.file_type == "txt"
    assert saved.status == "processing"
    assert saved.text_preview is None


def test_extract_text_reads_docx_paragraphs(tmp_path):
    from docx import Document as DocxDocument

    file_path = tmp_path / "sample.docx"
    document = DocxDocument()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.save(file_path)

    assert extract_text(str(file_path), "docx") == "First paragraph\n\nSecond paragraph"
